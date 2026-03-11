#!/usr/bin/env python3
"""
Однократная замена в готовых DOCX-шаблонах: «Индивидуальный предприниматель {{customer_name}}» → «{{customer_preamble}}».
Запуск из корня проекта: python scripts/patch_docx_customer_preamble.py
"""

import sys
from pathlib import Path

project_root = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(project_root))

from docx import Document

from app.config import TEMPLATES_DIR
from app.services.document_template_service import (
    replace_literals_with_placeholders_in_document,
)

REPLACEMENTS = [
    (
        "Индивидуальный предприниматель {{customer_name}}, именуемый в дальнейшем \"Заказчик\"",
        "{{customer_preamble}}, именуемый в дальнейшем \"Заказчик\"",
    ),
    (
        "Индивидуальный предприниматель {{customer_name}}, действующий на основании",
        "{{customer_preamble}}, действующий на основании",
    ),
]


def main():
    for filename in ("accounting_contract_template.docx", "marketing_contract_template.docx"):
        path = TEMPLATES_DIR / filename
        if not path.is_file():
            print(f"Пропуск (файл не найден): {filename}")
            continue
        doc = Document(path)
        replace_literals_with_placeholders_in_document(doc, REPLACEMENTS)
        doc.save(path)
        print(f"Обновлён: {filename}")


if __name__ == "__main__":
    main()
